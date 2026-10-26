"""Build anomalies.xlsx from the six exports.

Usage:
    python -m nfe.report_xlsx <jobs.csv> <qbo_txns.csv> --labor <labor.csv> \\
        --simpro-payments <sp.csv> --qbo-payments <qp.csv> \\
        [--user-log <ul.csv>] [--audit-log <al.csv>] [--snapshots DIR] \\
        -o anomalies.xlsx

Output is byte-identical across runs on the same inputs (fixed workbook
metadata + normalized zip timestamps). scripts/recalc.py independently
re-derives every computed figure from the workbook itself.
"""

from __future__ import annotations

import argparse
import sys

from openpyxl import Workbook

from . import config, snapshot
from .ingest import build_db, load_csv
from .rules import run_rules, mutation_exceptions
from .util import normalize_zip_in_place

SEV_ORDER = {"high": 0, "medium": 1, "low": 2}
BUCKET_ORDER = {config.BUCKET_PROBABLE: 0, config.BUCKET_UNRESOLVED: 1}

README_LINES = [
    "NFE anomaly report — Nexus Communications Technology",
    "",
    "EVERY ROW IN THIS WORKBOOK IS AN EXCEPTION REQUIRING INVESTIGATION,",
    "NEVER A FINDING OF WRONGDOING.",
    "",
    "What this report does NOT establish:",
    "- It does not establish intent, theft, or any criminal or civil liability.",
    "- It does not establish that any person named in a log entry acted",
    "  improperly. Shared logins, cleanup of duplicate jobs, re-quotes,",
    "  sync failures, timing lag and honest error are legitimate",
    "  explanations for every pattern flagged here.",
    "- Confidence is capped at 'medium' by the tool. 'High' requires a human",
    "  to test normal explanations against corroborating evidence",
    "  (two-source rule).",
    "",
    "Buckets: probable_exposure = corroborated by >= 2 independent sources;",
    "unresolved_exposure = single-source, needs more evidence.",
    "Labor cost benchmark: $104.00/hour fully loaded.",
    "",
    "Sheets: Exceptions (full register), Jobs_vs_Invoices, Labor,",
    "Payments, Exposure_Summary (ranked, strongest-evidenced first).",
]


def usd(cents) -> float:
    return round((cents or 0) / 100.0, 2)


def rank_key(e: dict):
    return (BUCKET_ORDER[e["bucket"]], SEV_ORDER[e["severity"]],
            -e["amount_cents"], e["exception_id"])


def build_workbook(conn, exceptions: list[dict]) -> Workbook:
    wb = Workbook()
    wb.properties.created = config.FIXED_DOC_TIMESTAMP
    wb.properties.modified = config.FIXED_DOC_TIMESTAMP
    wb.properties.creator = "nfe"
    wb.properties.lastModifiedBy = "nfe"

    ws = wb.active
    ws.title = "README"
    for line in README_LINES:
        ws.append([line])

    # --- Exceptions register ------------------------------------------------
    ws = wb.create_sheet("Exceptions")
    ws.append(["exception_id", "rule", "severity", "confidence", "bucket",
               "special_flag", "job_no", "customer", "amount_usd", "detail",
               "sources"])
    for e in sorted(exceptions, key=lambda e: e["exception_id"]):
        ws.append([e["exception_id"], e["rule"], e["severity"], e["confidence"],
                   e["bucket"], e["special_flag"], e["job_no"], e["customer"],
                   usd(e["amount_cents"]), e["detail"], e["sources"]])

    # --- Jobs vs Invoices -----------------------------------------------------
    flags_by_job: dict[str, set] = {}
    rules_by_job: dict[str, set] = {}
    for e in exceptions:
        if e["job_no"]:
            rules_by_job.setdefault(e["job_no"], set()).add(e["rule"])
            if e["special_flag"]:
                flags_by_job.setdefault(e["job_no"], set()).add(e["special_flag"])

    ws = wb.create_sheet("Jobs_vs_Invoices")
    ws.append(["job_no", "customer", "status", "job_value_usd",
               "qbo_invoiced_usd", "variance_usd", "qbo_invoice_count",
               "special_flags", "rules_hit"])
    for row in conn.execute("""
        SELECT j.job_no, j.customer, j.status, COALESCE(j.job_value_cents,0),
               COALESCE((SELECT SUM(q.amount_cents) FROM qbo_txns q
                         WHERE q.job_no = j.job_no), 0),
               COALESCE((SELECT COUNT(*) FROM qbo_txns q
                         WHERE q.job_no = j.job_no), 0)
        FROM jobs j ORDER BY j.job_no"""):
        job_no, cust, status, val, inv, n = row
        ws.append([job_no, cust, status, usd(val), usd(inv), usd(val - inv), n,
                   ";".join(sorted(flags_by_job.get(job_no, ()))),
                   ";".join(sorted(rules_by_job.get(job_no, ())))])

    # --- Labor -----------------------------------------------------------------
    ws = wb.create_sheet("Labor")
    ws.append(["job_no", "customer", "hours", "labor_cost_usd_at_104",
               "qbo_invoiced_usd", "rules_hit"])
    for job_no, cust, hours, cost, inv in conn.execute("""
        SELECT l.job_no, COALESCE(j.customer,''), SUM(l.hours),
               SUM(l.labor_cost_cents),
               COALESCE((SELECT SUM(q.amount_cents) FROM qbo_txns q
                         WHERE q.job_no = l.job_no), 0)
        FROM labor l LEFT JOIN jobs j ON j.job_no = l.job_no
        GROUP BY l.job_no ORDER BY l.job_no"""):
        ws.append([job_no, cust, round(hours, 2), usd(cost), usd(inv),
                   ";".join(sorted(r for r in rules_by_job.get(job_no, ())
                                   if r.startswith(("R05", "R06"))))])

    # --- Payments ---------------------------------------------------------------
    ws = wb.create_sheet("Payments")
    ws.append(["job_no", "customer", "simpro_paid_usd", "qbo_paid_usd",
               "delta_usd", "rules_hit"])
    pays = {}
    for job_no, amt in conn.execute(
            "SELECT job_no, SUM(amount_cents) FROM simpro_payments "
            "WHERE job_no IS NOT NULL GROUP BY job_no"):
        pays.setdefault(job_no, [0, 0])[0] = amt
    for job_no, amt in conn.execute(
            "SELECT job_no, SUM(amount_cents) FROM qbo_payments "
            "WHERE job_no IS NOT NULL GROUP BY job_no"):
        pays.setdefault(job_no, [0, 0])[1] = amt
    for job_no in sorted(pays):
        s, q = pays[job_no]
        cust = conn.execute("SELECT customer FROM jobs WHERE job_no=?",
                            (job_no,)).fetchone()
        ws.append([job_no, cust[0] if cust else "", usd(s), usd(q), usd(s - q),
                   ";".join(sorted(r for r in rules_by_job.get(job_no, ())
                                   if r.startswith(("R07", "R08", "R09", "R10"))))])

    # --- Exposure summary, ranked strongest-evidenced first ---------------------
    ws = wb.create_sheet("Exposure_Summary")
    ws.append(["rank", "exception_id", "rule", "bucket", "severity", "job_no",
               "customer", "amount_usd"])
    ranked = sorted(exceptions, key=rank_key)
    for i, e in enumerate(ranked, 1):
        ws.append([i, e["exception_id"], e["rule"], e["bucket"], e["severity"],
                   e["job_no"], e["customer"], usd(e["amount_cents"])])
    ws.append([])
    tot_p = sum(e["amount_cents"] for e in exceptions
                if e["bucket"] == config.BUCKET_PROBABLE)
    tot_u = sum(e["amount_cents"] for e in exceptions
                if e["bucket"] == config.BUCKET_UNRESOLVED)
    per_job_max: dict[str, int] = {}
    no_job_total = 0
    for e in exceptions:
        if e["job_no"]:
            per_job_max[e["job_no"]] = max(per_job_max.get(e["job_no"], 0),
                                           e["amount_cents"])
        else:
            no_job_total += e["amount_cents"]
    dedup = sum(per_job_max.values()) + no_job_total
    ws.append(["TOTAL_PROBABLE_EXPOSURE_USD", "", "", "", "", "", "", usd(tot_p)])
    ws.append(["TOTAL_UNRESOLVED_EXPOSURE_USD", "", "", "", "", "", "", usd(tot_u)])
    ws.append(["TOTAL_ALL_RULES_USD_MAY_DOUBLE_COUNT", "", "", "", "", "", "",
               usd(tot_p + tot_u)])
    ws.append(["DEDUPED_PER_JOB_MAX_EXPOSURE_USD", "", "", "", "", "", "",
               usd(dedup)])
    return wb


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(prog="python -m nfe.report_xlsx")
    ap.add_argument("jobs", help="full simPRO job list CSV (export 3)")
    ap.add_argument("qbo_txns", help="QBO all-invoices CSV (export 6)")
    ap.add_argument("--labor", help="simPRO labor/timesheets CSV (export 4)")
    ap.add_argument("--simpro-payments", dest="sp", help="simPRO payments CSV")
    ap.add_argument("--qbo-payments", dest="qp", help="QBO payments CSV")
    ap.add_argument("--user-log", dest="ul", help="simPRO user log CSV (export 2)")
    ap.add_argument("--audit-log", dest="al", help="QBO audit log CSV (export 1)")
    ap.add_argument("--snapshots", help="snapshot dir to check for mutations")
    ap.add_argument("-o", "--out", required=True, help="output .xlsx path")
    args = ap.parse_args(argv)

    datasets = {
        "jobs": load_csv(args.jobs, "jobs"),
        "qbo_txns": load_csv(args.qbo_txns, "qbo_txns"),
    }
    if args.labor:
        datasets["labor"] = load_csv(args.labor, "labor")
    if args.sp:
        datasets["simpro_payments"] = load_csv(args.sp, "simpro_payments")
    if args.qp:
        datasets["qbo_payments"] = load_csv(args.qp, "qbo_payments")
    if args.ul:
        datasets["user_log"] = load_csv(args.ul, "user_log")
    if args.al:
        datasets["audit_log"] = load_csv(args.al, "audit_log")

    conn = build_db(datasets)
    exceptions = run_rules(conn)
    if args.snapshots:
        exceptions += mutation_exceptions(snapshot.mutations(args.snapshots))

    wb = build_workbook(conn, exceptions)
    wb.save(args.out)
    normalize_zip_in_place(args.out)

    by_rule: dict[str, int] = {}
    for e in exceptions:
        by_rule[e["rule"]] = by_rule.get(e["rule"], 0) + 1
    print(f"wrote {args.out}: {len(exceptions)} exceptions")
    for rule in sorted(by_rule):
        print(f"  {rule}: {by_rule[rule]}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
