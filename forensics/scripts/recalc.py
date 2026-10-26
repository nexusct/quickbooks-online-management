#!/usr/bin/env python3
"""Independent verification of anomalies.xlsx.

Re-derives every computed figure from the workbook alone (no access to the
source CSVs) and reports the error count. Run after every report build:

    python scripts/recalc.py anomalies.xlsx    # must print '0 errors'
"""

from __future__ import annotations

import sys

from openpyxl import load_workbook

BUCKETS = {"probable_exposure": 0, "unresolved_exposure": 1}
SEVERITIES = {"high": 0, "medium": 1, "low": 2}


def cents(v) -> int:
    return round(float(v or 0) * 100)


def rows_of(ws):
    it = ws.iter_rows(values_only=True)
    header = next(it)
    for raw in it:
        if raw is None or all(c in (None, "") for c in raw):
            continue
        yield dict(zip(header, raw))


def main(path: str) -> int:
    errors: list[str] = []
    wb = load_workbook(path, read_only=False, data_only=True)

    for sheet in ("README", "Exceptions", "Jobs_vs_Invoices", "Labor",
                  "Payments", "Exposure_Summary"):
        if sheet not in wb.sheetnames:
            errors.append(f"missing sheet: {sheet}")
    if errors:
        print(f"{len(errors)} errors")
        for e in errors:
            print(" -", e)
        return 1

    readme = "\n".join(str(r[0] or "") for r in
                       wb["README"].iter_rows(values_only=True))
    if "EXCEPTION REQUIRING INVESTIGATION" not in readme:
        errors.append("README lost the exception-not-finding framing")
    if "does NOT establish" not in readme:
        errors.append("README lost the 'What this does NOT establish' section")

    # --- Exceptions register ------------------------------------------------
    exceptions = list(rows_of(wb["Exceptions"]))
    by_id = {}
    for e in exceptions:
        eid = e["exception_id"]
        if eid in by_id:
            errors.append(f"duplicate exception_id {eid}")
        by_id[eid] = e
        if e["bucket"] not in BUCKETS:
            errors.append(f"{eid}: bad bucket {e['bucket']!r}")
        if e["severity"] not in SEVERITIES:
            errors.append(f"{eid}: bad severity {e['severity']!r}")
        if cents(e["amount_usd"]) < 0:
            errors.append(f"{eid}: negative amount")
        n_sources = len([s for s in str(e["sources"]).split(";") if s])
        if e["bucket"] == "probable_exposure" and n_sources < 2:
            errors.append(f"{eid}: probable_exposure with <2 sources")
        if e["confidence"] == "high":
            errors.append(f"{eid}: tool-assigned confidence must cap at medium")

    # --- Jobs_vs_Invoices: variance must recompute --------------------------
    for r in rows_of(wb["Jobs_vs_Invoices"]):
        want = cents(r["job_value_usd"]) - cents(r["qbo_invoiced_usd"])
        if cents(r["variance_usd"]) != want:
            errors.append(f"job {r['job_no']}: variance {r['variance_usd']} "
                          f"!= value - invoiced")

    # --- Payments: delta must recompute --------------------------------------
    for r in rows_of(wb["Payments"]):
        want = cents(r["simpro_paid_usd"]) - cents(r["qbo_paid_usd"])
        if cents(r["delta_usd"]) != want:
            errors.append(f"payments {r['job_no']}: delta mismatch")

    # --- Exposure summary: ranking + footing ----------------------------------
    ranked, totals = [], {}
    for raw in wb["Exposure_Summary"].iter_rows(min_row=2, values_only=True):
        if raw is None or all(c in (None, "") for c in raw):
            continue
        first = str(raw[0] or "")
        if first.startswith(("TOTAL_", "DEDUPED_")):
            totals[first] = cents(raw[7])
        else:
            ranked.append(raw)

    if len(ranked) != len(exceptions):
        errors.append(f"exposure rows ({len(ranked)}) != exceptions "
                      f"({len(exceptions)})")
    prev_key = None
    for i, row in enumerate(ranked, 1):
        rank, eid, rule, bucket, sev, job_no, cust, amount = row
        if rank != i:
            errors.append(f"exposure rank {rank} at position {i}")
        src = by_id.get(eid)
        if src is None:
            errors.append(f"exposure row {eid} not in Exceptions sheet")
            continue
        if cents(amount) != cents(src["amount_usd"]) or bucket != src["bucket"] \
                or sev != src["severity"]:
            errors.append(f"exposure row {eid} disagrees with Exceptions sheet")
        key = (BUCKETS.get(bucket, 9), SEVERITIES.get(sev, 9),
               -cents(amount), str(eid))
        if prev_key is not None and key < prev_key:
            errors.append(f"exposure ordering broken at {eid} "
                          "(strongest-evidenced-first violated)")
        prev_key = key

    want_p = sum(cents(e["amount_usd"]) for e in exceptions
                 if e["bucket"] == "probable_exposure")
    want_u = sum(cents(e["amount_usd"]) for e in exceptions
                 if e["bucket"] == "unresolved_exposure")
    per_job_max: dict[str, int] = {}
    no_job = 0
    for e in exceptions:
        job = str(e["job_no"] or "")
        amt = cents(e["amount_usd"])
        if job:
            per_job_max[job] = max(per_job_max.get(job, 0), amt)
        else:
            no_job += amt
    checks = {
        "TOTAL_PROBABLE_EXPOSURE_USD": want_p,
        "TOTAL_UNRESOLVED_EXPOSURE_USD": want_u,
        "TOTAL_ALL_RULES_USD_MAY_DOUBLE_COUNT": want_p + want_u,
        "DEDUPED_PER_JOB_MAX_EXPOSURE_USD": sum(per_job_max.values()) + no_job,
    }
    for label, want in checks.items():
        if label not in totals:
            errors.append(f"missing total row {label}")
        elif totals[label] != want:
            errors.append(f"{label} does not foot: sheet {totals[label]} "
                          f"recomputed {want}")

    print(f"{len(errors)} errors")
    for e in errors:
        print(" -", e)
    return 1 if errors else 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("usage: python scripts/recalc.py anomalies.xlsx")
        sys.exit(2)
    sys.exit(main(sys.argv[1]))
